#!/usr/bin/env python3
import json, sys, os
from openpyxl import Workbook, load_workbook

def handle_read(args):
    path = args["path"]
    sheet_name = args.get("sheet", "")
    max_rows = args.get("max_rows", 50)
    ext = os.path.splitext(path)[1].lower()

    if ext == ".csv":
        import csv
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            rows = list(reader)
        headers = rows[0] if rows else []
        data = rows[1:1+max_rows]
        text = " | ".join(headers) + "\n" + "-" * 40 + "\n"
        for r in data:
            text += " | ".join(r) + "\n"
        text += f"\n({len(rows)-1} rows total, showing {len(data)})"
        return {"text": text}

    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name:
        ws = wb[sheet_name]
    else:
        ws = wb.active
        sheet_name = ws.title

    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    text = f"Sheet: {sheet_name}\n"
    text += " | ".join(str(h or "") for h in headers) + "\n" + "-" * 40 + "\n"
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if count >= max_rows: break
        text += " | ".join(str(c or "") for c in row) + "\n"
        count += 1
    text += f"\n({count}+ rows total, showing {count})"
    wb.close()
    return {"text": text}

def handle_write(args):
    path = args["path"]
    headers = args["headers"]
    rows = args["rows"]
    ext = os.path.splitext(path)[1].lower()

    if ext == ".csv":
        import csv
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(headers)
            for r in rows: w.writerow(r)
        return {"path": path}

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows: ws.append(r)
    wb.save(path)
    return {"path": path}

def handle_transform(args):
    path = args["path"]
    ops = args.get("operations", [])
    out_path = args.get("output_path", path)

    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        import csv
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            rows = list(reader)
        headers, data = rows[0], rows[1:]
    else:
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        data = [list(row) for row in ws.iter_rows(min_row=2, values_only=True)]
        wb.close()

    for op in ops:
        op_type = op.get("type", "")
        if op_type == "filter":
            col = op["column"]
            val = op["value"]
            idx = headers.index(col) if col in headers else -1
            if idx >= 0:
                data = [r for r in data if idx < len(r) and str(r[idx]) == str(val)]
        elif op_type == "sort":
            col = op["column"]
            desc = op.get("desc", False)
            idx = headers.index(col) if col in headers else -1
            if idx >= 0:
                data.sort(key=lambda r: str(r[idx]) if idx < len(r) else "", reverse=desc)
        elif op_type == "add_column":
            headers.append(op["name"])
            default = op.get("default", "")
            for r in data: r.append(default)

    out_ext = os.path.splitext(out_path)[1].lower()
    if out_ext == ".csv":
        import csv
        with open(out_path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(headers)
            for r in data: w.writerow(r)
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for r in data: ws.append(r)
        wb.save(out_path)

    return {"text": f"✅ Transformed: {len(headers)} cols, {len(data)} rows", "output": out_path}

def handle_errors(args):
    path = args["path"]
    ext = os.path.splitext(path)[1].lower()
    issues = []

    if ext == ".csv":
        import csv
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            rows = list(reader)
        headers, data = rows[0], rows[1:]
    else:
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        data = [list(row) for row in ws.iter_rows(min_row=2, values_only=True)]
        wb.close()

    col_count = len(headers)
    for i, r in enumerate(data):
        if len(r) != col_count:
            issues.append(f"Row {i+2}: expected {col_count} cols, got {len(r)}")
        for j, c in enumerate(r):
            if c is None:
                issues.append(f"Row {i+2}, col {headers[j]}: empty cell")
    for j, h in enumerate(headers):
        col_vals = [str(r[j]) for r in data if j < len(r) and r[j] is not None]
        if len(set(col_vals)) == 1 and len(col_vals) > 1:
            issues.append(f"Col '{h}': all values identical ({col_vals[0]})")

    text = f"Checked {len(headers)} cols x {len(data)} rows\n"
    if issues:
        text += "\nIssues:\n" + "\n".join(issues[:20])
        if len(issues) > 20: text += f"\n... and {len(issues)-20} more"
    else:
        text += "\n✅ No issues found"
    return {"text": text}

def main():
    payload = json.loads(sys.stdin.read())
    cmd, args = payload["cmd"], payload.get("args", {})
    handlers = {"read": handle_read, "write": handle_write, "transform": handle_transform, "errors": handle_errors}
    h = handlers.get(cmd)
    if not h: print(json.dumps({"error": f"Unknown: {cmd}"})); sys.exit(1)
    try: print(json.dumps(h(args)))
    except Exception as e: print(json.dumps({"error": str(e)})); sys.exit(1)

if __name__ == "__main__":
    main()
